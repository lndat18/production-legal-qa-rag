"""Đọc `.docx` bằng python-docx và trích xuất Block theo đúng thứ tự.

Duyệt ``document.element.body`` chứ không duyệt ``doc.paragraphs`` rồi
``doc.tables`` riêng — cách sau làm mất thứ tự xen kẽ giữa đoạn văn và bảng.
Style/bold/nghiêng của đoạn văn được giữ lại như tín hiệu phụ: rule nhận
diện heading (mục 3 spec) không bao giờ đọc các giá trị này, vì style trong
corpus thực tế không nhất quán; tín hiệu này chỉ dùng để serialize block cho
Gemini khi chuyển đổi front matter/back matter (``serialize_blocks_for_llm``,
mục 1.1 spec).
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from production_legal_qa_rag.formatting.patterns import normalize_text
from production_legal_qa_rag.formatting.tables import table_to_markdown


@dataclass(frozen=True, slots=True)
class Block:
    """Một phần tử nội dung lấy từ DOCX, giữ đúng thứ tự xuất hiện."""

    kind: Literal["paragraph", "table"]
    text: str  # với table: markdown (hoặc HTML) đã render sẵn
    style: str | None = None  # tên style gốc, chỉ là tín hiệu phụ
    is_bold: bool = False
    is_italic: bool = False


def _paragraph_all_runs_have(paragraph: Paragraph, attribute: str) -> bool:
    """Mọi run (có chữ) của đoạn có bật thuộc tính ``attribute`` hay không.

    Dùng chung cho bold/italic. Chỉ là tín hiệu phụ. Rule nhận diện heading
    không bao giờ được đọc giá trị này: phần lớn paragraph trong corpus có
    style "Normal" và bold/italic không phủ đủ toàn đoạn.
    """
    runs = [run for run in paragraph.runs if run.text.strip()]
    return bool(runs) and all(getattr(run, attribute) for run in runs)


def _paragraph_is_bold(paragraph: Paragraph) -> bool:
    """Đoạn có in đậm toàn bộ hay không."""
    return _paragraph_all_runs_have(paragraph, "bold")


def _paragraph_is_italic(paragraph: Paragraph) -> bool:
    """Đoạn có in nghiêng toàn bộ hay không."""
    return _paragraph_all_runs_have(paragraph, "italic")


def read_docx(path: str | Path) -> list[Block]:
    """Đọc DOCX, giữ đúng thứ tự xen kẽ giữa paragraph và table.

    Args:
        path: Đường dẫn tới file `.docx`.

    Returns:
        Danh sách Block theo đúng thứ tự xuất hiện trong tài liệu gốc.
    """
    document = Document(str(path))
    blocks: list[Block] = []

    for element in document.element.body.iterchildren():
        if isinstance(element, CT_P):
            paragraph = Paragraph(element, document)
            text = normalize_text(paragraph.text)
            if not text:
                continue
            style = paragraph.style.name if paragraph.style is not None else None
            blocks.append(
                Block(
                    kind="paragraph",
                    text=text,
                    style=style,
                    is_bold=_paragraph_is_bold(paragraph),
                    is_italic=_paragraph_is_italic(paragraph),
                )
            )
        elif isinstance(element, CT_Tbl):
            markdown_table = table_to_markdown(Table(element, document))
            if markdown_table:
                blocks.append(Block(kind="table", text=markdown_table))

    return blocks


def serialize_blocks_for_llm(blocks: list[Block]) -> str:
    """Render Block thành text kèm tín hiệu bold/nghiêng, dùng làm input cho Gemini.

    Paragraph in đậm/nghiêng toàn dòng được bọc sẵn ``**``/``*`` (cả hai cùng
    lúc thì bọc ``***``) — Gemini chỉ cần giữ nguyên định dạng khi chuyển đổi
    (mục 1.1 spec), không phải tự đoán từ văn bản thô. Bảng đã có sẵn
    markdown/HTML từ ``tables.table_to_markdown``, giữ nguyên văn.

    Args:
        blocks: Block front matter hoặc back matter, theo đúng thứ tự gốc.

    Returns:
        Text đã ghép, mỗi block cách nhau một dòng trống.
    """
    lines: list[str] = []
    for block in blocks:
        if block.kind == "table":
            lines.append(block.text)
            continue
        text = block.text
        if block.is_bold and block.is_italic:
            text = f"***{text}***"
        elif block.is_bold:
            text = f"**{text}**"
        elif block.is_italic:
            text = f"*{text}*"
        lines.append(text)
    return "\n\n".join(lines)
