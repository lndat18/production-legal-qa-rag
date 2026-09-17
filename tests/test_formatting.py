"""Bộ test cho bước formatting (DOCX -> Markdown), theo `formatting_spec.md`.

Thiết kế mới (mục 1.1, 1.3): front matter/back matter không còn trích
field/YAML, chỉ xác định biên deterministic rồi chia chunk (mục 1.2), dựng
prompt thuần (`frontmatter.build_prompts`/`backmatter.build_prompts`, không
I/O), gọi Groq đúng 1 lần cho toàn bộ file qua
`llm_client.convert_chunks_concurrently` (điều phối bởi `pipeline.py`, chạy
đồng thời 2 worker nếu có `GROQ_API_KEY_2`), rồi ghép kết quả thuần
(`frontmatter.assemble`/`backmatter.assemble`). Test suite này KHÔNG BAO GIỜ
gọi Groq API thật (fixture `_default_llm_stub` mock
`llm_client.convert_chunks_concurrently` cho toàn bộ session, mặc định trả
`None` cho mọi prompt) -- không nên phụ thuộc mạng/API key khi chạy CI.

Cấu trúc theo từng module: models, patterns (bao gồm bug setext heading, mục
1.1), docx_reader, tables, emitter, validator, frontmatter, backmatter,
llm_client (bao gồm dispatch 2 key, mục 1.3), pipeline (integration trên
`data/raw/*.docx` thật + CLI).
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path
from typing import get_args
from unittest.mock import Mock

import pytest
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pydantic import ValidationError
from typer.testing import CliRunner

from production_legal_qa_rag.config import LLMSettings
from production_legal_qa_rag.formatting import (
    backmatter,
    emitter,
    frontmatter,
    llm_client,
    pipeline,
    tables,
    validator,
)
from production_legal_qa_rag.formatting.docx_reader import (
    Block,
    chunk_blocks_for_llm,
    read_docx,
    serialize_blocks_for_llm,
)
from production_legal_qa_rag.formatting.models import (
    FormattingResult,
    QcWarning,
    QcWarningCode,
)
from production_legal_qa_rag.formatting.patterns import (
    RE_DIEM,
    RE_DIEU,
    RE_FOOTNOTE_MARKER,
    RE_KHOAN,
    escape_setext_underline,
    is_structural,
    normalize_text,
    sort_key,
    strip_markers,
)
from production_legal_qa_rag.formatting.pipeline import (
    convert_directory,
    convert_docx_to_markdown,
    write_atomic,
)

PROJECT_ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = PROJECT_ROOT / "data" / "raw"
REFERENCE_MD_DIR = PROJECT_ROOT / "data" / "markdown"

# Rich (dùng bởi Typer để render --help) vẫn chèn mã CSI cho style (bold, dim)
# dù đã set NO_COLOR -- NO_COLOR chỉ tắt màu, không tắt style. Gỡ mã ANSI trước
# khi so khớp chuỗi để không phụ thuộc vào việc Rich style output ra sao.
_ANSI_ESCAPE_RE = re.compile(r"\x1b\[[0-9;]*[A-Za-z]")

RAW_FILES = sorted(RAW_DIR.glob("*.docx")) if RAW_DIR.exists() else []

# Toàn bộ mã QcWarning hợp lệ đọc trực tiếp từ `QcWarningCode` (models.py) --
# không hard-code lại danh sách ở đây, tránh lệch nhau khi thêm rule QC mới.
KNOWN_WARNING_CODES = set(get_args(QcWarningCode))


def P(text: str, *, style: str | None = None, bold: bool = False) -> Block:
    return Block(kind="paragraph", text=text, style=style, is_bold=bold)


def T(text: str) -> Block:
    return Block(kind="table", text=text)


# ==========================================================================
# Không bao giờ gọi Groq API thật trong test suite (CI không có
# GROQ_API_KEY; máy dev có thể có key thật -- không nên phụ thuộc vào việc
# thiếu key mới an toàn, mock hẳn ở mức `llm_client.convert_chunks_concurrently`,
# điểm gọi Groq DUY NHẤT của `pipeline.py` từ mục 1.3 trở đi).
# ==========================================================================

# Giữ tham chiếu tới hàm THẬT trước khi fixture dưới đây ghi đè
# `llm_client.convert_chunks_concurrently` -- các test của "llm_client.py"
# cần gọi đúng implementation thật (chỉ mock `_client`/`_client_2`, không
# mock `convert_chunks_concurrently` chính nó) để kiểm tra logic
# dispatch/thread/thứ tự thật của nó.
_REAL_CONVERT_CHUNKS_CONCURRENTLY = llm_client.convert_chunks_concurrently

# Giữ tham chiếu tới object `lru_cache` THẬT của các singleton module-level
# -- một số test monkeypatch trực tiếp `llm_client._client`/`_client_2` bằng
# 1 lambda trần (không có `.cache_clear`), nên fixture dọn dẹp dưới đây phải
# luôn thao tác trên object gốc này, không đọc lại qua `llm_client._client`
# (có thể đang bị monkeypatch tại thời điểm teardown chạy, thứ tự teardown
# giữa các fixture không đảm bảo monkeypatch đã revert trước).
_REAL_CLIENT_FACTORY = llm_client._client
_REAL_CLIENT_2_FACTORY = llm_client._client_2
_REAL_RATE_LIMITER_FACTORY = llm_client._rate_limiter
_REAL_RATE_LIMITER_2_FACTORY = llm_client._rate_limiter_2


@pytest.fixture(autouse=True)
def _default_llm_stub(monkeypatch: pytest.MonkeyPatch):
    """Mặc định `llm_client.convert_chunks_concurrently` trả `None` cho MỌI
    prompt (Groq "lỗi") ở MỌI test, trừ khi test tự monkeypatch lại giá trị
    khác bên trong. Test cần gọi implementation THẬT dùng
    `_REAL_CONVERT_CHUNKS_CONCURRENTLY` (tham chiếu lưu trước khi fixture
    này chạy) để không bị chính stub này che mất.
    """
    monkeypatch.setattr(
        llm_client,
        "convert_chunks_concurrently",
        lambda prompts: [None] * len(prompts),
    )


@pytest.fixture(autouse=True)
def _reset_llm_client_singletons():
    """Dọn state của các singleton `lru_cache` module-level trong
    `llm_client.py` (`_client`/`_client_2`/`_rate_limiter`/`_rate_limiter_2`)
    trước/sau mỗi test -- không dọn thì entry của sliding-window rate limiter
    hay client giả ghi nhận ở 1 test có thể rò rỉ sang test kế tiếp trong
    cùng phiên pytest, gây `time.sleep` thật ngoài ý muốn khi test không tự
    mock `time` (mục 1.2 spec), hoặc dùng nhầm client giả của test trước.
    """
    _REAL_CLIENT_FACTORY.cache_clear()
    _REAL_CLIENT_2_FACTORY.cache_clear()
    _REAL_RATE_LIMITER_FACTORY.cache_clear()
    _REAL_RATE_LIMITER_2_FACTORY.cache_clear()
    yield
    _REAL_CLIENT_FACTORY.cache_clear()
    _REAL_CLIENT_2_FACTORY.cache_clear()
    _REAL_RATE_LIMITER_FACTORY.cache_clear()
    _REAL_RATE_LIMITER_2_FACTORY.cache_clear()


def _forbid_llm_calls(monkeypatch: pytest.MonkeyPatch) -> None:
    """Fail nếu `convert_chunks_concurrently` được gọi với ÍT NHẤT 1 prompt
    thật. Gọi với danh sách RỖNG vẫn hợp lệ -- `pipeline.py` luôn gọi hàm
    này đúng 1 lần dù front/back matter rỗng (mục 1.3 spec: nối
    `before_prompts + after_prompts + chunk_prompts`, có thể rỗng cả 3),
    `convert_chunks_concurrently([])` trả `[]` ngay mà không chạm mạng/API
    key -- đây không phải "gọi Groq" theo nghĩa cần cấm.
    """

    def _fail(prompts: list[str]) -> list[str | None]:
        if prompts:
            raise AssertionError(
                "llm_client.convert_chunks_concurrently không được gọi với "
                "prompt thật ở đây"
            )
        return []

    monkeypatch.setattr(llm_client, "convert_chunks_concurrently", _fail)


def _stub_convert_chunks_concurrently(
    monkeypatch: pytest.MonkeyPatch, values: list[str | None]
) -> list[str]:
    """Trả về đúng `values` cho lần gọi `convert_chunks_concurrently` DUY
    NHẤT của `pipeline.convert_docx_to_markdown` (mục 1.3 spec -- gộp toàn
    bộ job front/back matter của 1 file thành 1 lệnh gọi). Trả về danh sách
    prompt đã nhận được (để test kiểm tra nội dung/số lượng), assert ngay số
    lượng prompt khớp `values` để test fail rõ ràng nếu ranh giới cắt sai ở
    `pipeline.py`.
    """
    captured: list[str] = []

    def fake(prompts: list[str]) -> list[str | None]:
        captured.extend(prompts)
        assert len(prompts) == len(values), (
            f"kỳ vọng {len(values)} prompt nhưng nhận {len(prompts)}: {prompts}"
        )
        return values

    monkeypatch.setattr(llm_client, "convert_chunks_concurrently", fake)
    return captured


# ==========================================================================
# models.py -- schema
# ==========================================================================


def test_qc_warning_default_detail_rong():
    warning = QcWarning(code="no_heading")
    assert warning.detail == ""


def test_qc_warning_code_khong_hop_le_bi_tu_choi():
    with pytest.raises(ValidationError):
        QcWarning(code="ma_khong_ton_tai")  # type: ignore[arg-type]


def test_formatting_result_warnings_mac_dinh_rong():
    result = FormattingResult(markdown="x")
    assert result.warnings == []


def test_formatting_result_thieu_markdown_bao_loi():
    with pytest.raises(ValidationError):
        FormattingResult()  # type: ignore[call-arg]


def test_qc_warning_code_gom_du_ba_ma_moi():
    assert "llm_frontmatter_conversion_failed" in KNOWN_WARNING_CODES
    assert "llm_backmatter_conversion_failed" in KNOWN_WARNING_CODES
    assert "frontmatter_title_not_found" in KNOWN_WARNING_CODES


# ==========================================================================
# patterns.py
# ==========================================================================


def test_normalize_text_gop_khoang_trang_va_nbsp():
    assert normalize_text("Điều 1.  Phạm  vi") == "Điều 1. Phạm vi"


def test_normalize_text_strip_dau_cuoi():
    assert normalize_text("  Chương I  ") == "Chương I"


@pytest.mark.parametrize(
    ("smaller", "larger"),
    [
        ("48", "48a"),
        ("48a", "48b"),
        ("48b", "49"),
        ("1", "1a"),
        ("9", "10"),
    ],
)
def test_sort_key_thu_tu_so_hieu_co_hau_to_chu(smaller, larger):
    assert sort_key(smaller) < sort_key(larger)


def test_re_dieu_ho_tro_so_hieu_chu():
    match = RE_DIEU.match("Điều 41a. Điều khoản chuyển tiếp")
    assert match is not None
    assert match.group(1) == "41a"
    assert match.group(2) == "Điều khoản chuyển tiếp"


def test_re_dieu_khong_khop_khi_thieu_dau_cham():
    # "Điều 2 của Luật số 46/2014/QH13 quy định..." (trích dẫn) không được
    # nhận nhầm thành heading Điều.
    assert RE_DIEU.match("Điều 2 của Luật số 46/2014/QH13 quy định như sau:") is None


def test_re_khoan_khop_va_giu_noi_dung():
    match = RE_KHOAN.match("1. Người lao động làm việc theo hợp đồng.")
    assert match is not None
    assert match.group(1) == "1"
    assert match.group(2) == "Người lao động làm việc theo hợp đồng."


def test_re_diem_hai_dang_chu_va_gach_dau_dong():
    letter = RE_DIEM.match("a) Doanh nghiệp theo quy định.")
    dash = RE_DIEM.match("- Một mục không đánh số.")
    assert letter is not None and letter.group(1) == "a"
    assert dash is not None and dash.group(1) is None


def test_is_structural_nhan_dien_dieu_va_chuong_la_cau_truc():
    assert is_structural("Điều 5. Tên điều")
    assert is_structural("Chương IV")
    assert not is_structural("Người lao động làm việc theo hợp đồng.")


def test_is_structural_khong_nhan_dien_khoan():
    # Khoản KHÔNG được coi là biên front matter -- các dòng "1. Luật Nhà giáo
    # số 73/2025/QH15..." trong dẫn nhập văn bản hợp nhất không được nhận
    # nhầm thành biên (mục 1.1 spec).
    assert not is_structural("1. Luật Nhà giáo số 73/2025/QH15.")


# ==========================================================================
# patterns.py -- strip_markers()/RE_FOOTNOTE_MARKER, marker chú thích "[n]"
# dính liền trong vùng nội dung ở giữa (formatting_spec.md mục 1.1, "Lưu ý
# quan trọng"). Regression cho lỗi footnote marker phá RE_KHOAN/RE_DIEU.
# ==========================================================================


def test_strip_markers_go_marker_dinh_ngay_sau_so_khoan():
    # "1.[2] Bảo hiểm..." -- không strip trước thì RE_KHOAN (yêu cầu khoảng
    # trắng ngay sau dấu chấm) sẽ không khớp và mất heading Khoản.
    assert (
        strip_markers("1.[2] Bảo hiểm y tế là hình thức bắt buộc.")
        == "1. Bảo hiểm y tế là hình thức bắt buộc."
    )


def test_strip_markers_go_marker_o_cuoi_dong():
    assert strip_markers("Điều 7a. Bảo hiểm Xã hội[16]") == "Điều 7a. Bảo hiểm Xã hội"


def test_strip_markers_khong_nuot_so_hieu_khi_co_dau_cham_o_giua():
    # "10.[15]" -- chữ số trước "[" không liền kề (có dấu chấm chen giữa) nên
    # KHÔNG bị coi là chữ số lặp cần nuốt, số hiệu "10." phải giữ nguyên.
    assert strip_markers("10.[15] Điều khoản") == "10. Điều khoản"


def test_strip_markers_bo_chu_so_lap_ngay_truoc_marker():
    # "3.3[3]" -- chữ số lặp dính liền ngay trước "[" và BẰNG số marker, bị
    # nuốt cùng marker, chỉ giữ lại "3." gốc.
    assert strip_markers("3.3[3] Nội dung.") == "3. Nội dung."


def test_strip_markers_va_lap_lai_khoang_trang_thieu_sau_marker():
    # Sau khi gỡ marker dính liền, "a)[3]Thành lập" mất khoảng trắng --
    # RE_MISSING_SPACE phải vá lại.
    assert strip_markers("a)[3]Thành lập") == "a) Thành lập"


def test_strip_markers_don_dep_khoang_trang_thua_truoc_dau_cau():
    assert strip_markers("Nội dung[5] , tiếp theo.") == "Nội dung, tiếp theo."


def test_strip_markers_khong_co_marker_giu_nguyen():
    assert strip_markers("Không có marker nào ở đây.") == "Không có marker nào ở đây."


def test_re_footnote_marker_khop_va_lay_duoc_so():
    match = RE_FOOTNOTE_MARKER.search("Bảo hiểm Xã hội[16]")
    assert match is not None
    assert match.group(1) == "16"


# ==========================================================================
# patterns.py -- escape_setext_underline (formatting_spec.md mục 1.1, "[MỚI
# 2026-09-17]" -- bug setext heading phát hiện trên `Luật bảo hiểm xã hội.md`).
# ==========================================================================


def test_escape_setext_underline_da_co_dong_trong_khong_doi():
    markdown = "VĂN PHÒNG QUỐC HỘI\n\n--------"
    assert escape_setext_underline(markdown) == markdown


def test_escape_setext_underline_thieu_dong_trong_chen_them():
    markdown = "VĂN PHÒNG QUỐC HỘI\n--------"
    assert escape_setext_underline(markdown) == "VĂN PHÒNG QUỐC HỘI\n\n--------"


def test_escape_setext_underline_dau_bang_cung_duoc_xu_ly():
    markdown = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n===================="
    assert (
        escape_setext_underline(markdown)
        == "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n\n===================="
    )


def test_escape_setext_underline_dong_gach_ngang_dau_van_ban_giu_nguyen():
    # Dòng gạch ngang là dòng ĐẦU TIÊN (không có dòng nào phía trên để chèn
    # dòng trống vào giữa) -- giữ nguyên.
    markdown = "--------\n\nNội dung."
    assert escape_setext_underline(markdown) == markdown


def test_escape_setext_underline_khong_du_3_ky_tu_khong_phai_setext():
    # "--" chỉ 2 ký tự -- dưới ngưỡng CommonMark, không khớp
    # RE_SETEXT_UNDERLINE, giữ nguyên dù đứng liền kề dòng text.
    markdown = "Text\n--"
    assert escape_setext_underline(markdown) == markdown


def test_escape_setext_underline_khong_dung_vao_dong_ke_bang():
    # Dòng phân cách bảng "| --- | --- |" KHÔNG khớp RE_SETEXT_UNDERLINE (có
    # ký tự "|"/khoảng trắng xen giữa, không phải thuần "-"/"=").
    markdown = "Tiêu đề\n| --- | --- |"
    assert escape_setext_underline(markdown) == markdown


def test_escape_setext_underline_nhieu_cho_trong_cung_mot_van_ban():
    markdown = "A\n---\n\nB\n===\n\nC"
    assert escape_setext_underline(markdown) == "A\n\n---\n\nB\n\n===\n\nC"


# ==========================================================================
# docx_reader.py -- đọc DOCX thật, giữ thứ tự paragraph/table xen kẽ
# ==========================================================================


def test_read_docx_giu_dung_thu_tu_xen_ke(tmp_path: Path):
    document = Document()
    document.add_paragraph("Điều 1. Phạm vi điều chỉnh")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cột A"
    table.cell(0, 1).text = "Cột B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    document.add_paragraph("1. Nội dung khoản một.")
    document.add_paragraph("   ")  # đoạn trắng, phải bị bỏ qua

    docx_path = tmp_path / "sample.docx"
    document.save(docx_path)

    blocks = read_docx(docx_path)
    assert [block.kind for block in blocks] == ["paragraph", "table", "paragraph"]
    assert blocks[0].text == "Điều 1. Phạm vi điều chỉnh"
    assert blocks[1].text.startswith("| Cột A | Cột B |")
    assert blocks[2].text == "1. Nội dung khoản một."


def test_read_docx_in_dam_toan_doan_la_bold(tmp_path: Path):
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("NGHỊ ĐỊNH")
    run.bold = True

    docx_path = tmp_path / "bold.docx"
    document.save(docx_path)

    blocks = read_docx(docx_path)
    assert blocks[0].is_bold is True
    assert blocks[0].is_italic is False


def test_read_docx_in_nghieng_toan_doan_la_italic(tmp_path: Path):
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Độc lập - Tự do - Hạnh phúc")
    run.italic = True

    docx_path = tmp_path / "italic.docx"
    document.save(docx_path)

    blocks = read_docx(docx_path)
    assert blocks[0].is_italic is True
    assert blocks[0].is_bold is False


def test_serialize_blocks_for_llm_bold_duoc_bao_bang_sao_kep():
    blocks = [P("CHÍNH PHỦ", bold=True)]
    assert serialize_blocks_for_llm(blocks) == "**CHÍNH PHỦ**"


def test_serialize_blocks_for_llm_italic_duoc_bao_bang_mot_sao():
    blocks = [Block(kind="paragraph", text="Độc lập - Tự do", is_italic=True)]
    assert serialize_blocks_for_llm(blocks) == "*Độc lập - Tự do*"


def test_serialize_blocks_for_llm_dam_va_nghieng_duoc_bao_bang_ba_sao():
    blocks = [Block(kind="paragraph", text="CHÍNH PHỦ", is_bold=True, is_italic=True)]
    assert serialize_blocks_for_llm(blocks) == "***CHÍNH PHỦ***"


def test_serialize_blocks_for_llm_khong_dinh_dang_giu_nguyen():
    blocks = [P("Số: 293/2025/NĐ-CP")]
    assert serialize_blocks_for_llm(blocks) == "Số: 293/2025/NĐ-CP"


def test_serialize_blocks_for_llm_nhieu_block_cach_nhau_dong_trong():
    blocks = [P("A"), P("B")]
    assert serialize_blocks_for_llm(blocks) == "A\n\nB"


def test_serialize_blocks_for_llm_giu_nguyen_bang():
    table_text = "| a | b |\n| --- | --- |\n| 1 | 2 |"
    blocks = [T(table_text)]
    assert serialize_blocks_for_llm(blocks) == table_text


# ==========================================================================
# docx_reader.py -- chunk_blocks_for_llm (formatting_spec.md mục 1.2)
# ==========================================================================


def test_chunk_blocks_for_llm_rong_tra_ve_danh_sach_rong():
    assert chunk_blocks_for_llm([], token_limit=1500) == []


def test_chunk_blocks_for_llm_duoi_gioi_han_mot_chunk_duy_nhat():
    blocks = [P("Đoạn ngắn 1."), P("Đoạn ngắn 2."), P("Đoạn ngắn 3.")]
    chunks = chunk_blocks_for_llm(blocks, token_limit=1500)
    assert chunks == [blocks]


def test_chunk_blocks_for_llm_vuot_gioi_han_tach_chunk_moi():
    # Mỗi block ~40 ký tự => ~16 token (heuristic len // 2.5). token_limit=20
    # chỉ đủ cho 1 block mỗi chunk.
    blocks = [P("A" * 40), P("B" * 40), P("C" * 40)]
    chunks = chunk_blocks_for_llm(blocks, token_limit=20)
    assert chunks == [[blocks[0]], [blocks[1]], [blocks[2]]]


def test_chunk_blocks_for_llm_khong_cat_giua_mot_block():
    # 1 block tự nó đã vượt token_limit vẫn phải nằm trọn trong 1 chunk
    # riêng -- không có cách nào chia nhỏ hơn mà không cắt giữa block.
    huge_block = P("X" * 10_000)
    blocks = [P("nhỏ"), huge_block, P("nhỏ 2")]
    chunks = chunk_blocks_for_llm(blocks, token_limit=100)
    matching_chunks = [chunk for chunk in chunks if huge_block in chunk]
    assert len(matching_chunks) == 1
    assert matching_chunks[0] == [huge_block]
    # Mọi block gốc phải xuất hiện đúng 1 lần, đúng thứ tự, không bị cắt.
    flattened = [block for chunk in chunks for block in chunk]
    assert flattened == blocks


def test_chunk_blocks_for_llm_giu_dung_thu_tu_tong_hop():
    blocks = [P(f"block {i} " + "x" * 30) for i in range(5)]
    chunks = chunk_blocks_for_llm(blocks, token_limit=30)
    flattened = [block for chunk in chunks for block in chunk]
    assert flattened == blocks
    assert len(chunks) > 1


@pytest.mark.skipif(not RAW_FILES, reason="data/raw/*.docx không tồn tại")
def test_chunk_blocks_for_llm_tren_bao_hiem_y_te_that_tach_nhieu_chunk():
    """Mục 7 spec: `Luật bảo hiểm y tế.docx` là ca kiểm thử chính cho việc
    chunk hoạt động đúng trên corpus thật -- back matter ~11.671 token ước
    lượng, vượt xa `CHUNK_TOKEN_LIMIT=1500` nếu gửi nguyên khối, nên bắt
    buộc phải tách thành nhiều chunk, mỗi chunk dưới ngưỡng. Test này KHÔNG
    gọi Groq -- xác định biên (`frontmatter.find_boundary`,
    `backmatter.split_backmatter`) và `chunk_blocks_for_llm` đều là hàm
    thuần, cục bộ, không phụ thuộc mạng/API key."""
    path = RAW_DIR / "Luật bảo hiểm y tế.docx"
    if not path.exists():
        pytest.skip("data/raw/Luật bảo hiểm y tế.docx không tồn tại")

    blocks = read_docx(path)
    fm_boundary = frontmatter.find_boundary(blocks)
    rest = blocks[fm_boundary:]
    _, back_blocks, _ = backmatter.split_backmatter(rest)

    assert back_blocks, "kỳ vọng file này có back matter (mục 1.2 spec)"

    chunk_token_limit = LLMSettings.model_fields["chunk_token_limit"].default
    chunks = chunk_blocks_for_llm(back_blocks, chunk_token_limit)

    assert len(chunks) > 1
    # Không có block đơn lẻ nào trên corpus thật đủ lớn để tự vượt ngưỡng
    # (mục 1.2 "Không bao giờ cắt giữa 1 block" chỉ là ngoại lệ lý thuyết) --
    # nên mọi chunk thực tế phải nằm dưới `chunk_token_limit`.
    for chunk in chunks:
        estimated_tokens = sum(len(block.text) for block in chunk) / 2.5
        assert estimated_tokens <= chunk_token_limit
    # Không cắt giữa block, giữ đúng thứ tự gốc.
    flattened = [block for chunk in chunks for block in chunk]
    assert flattened == back_blocks


# ==========================================================================
# tables.py
# ==========================================================================


def _make_docx_table(rows: list[list[str]]):
    document = Document()
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = value
    return table


def test_table_to_markdown_bang_du_lieu_ra_pipe_table():
    table = _make_docx_table([["Vùng", "Mức lương"], ["I", "5.310.000"]])
    rendered = tables.table_to_markdown(table)
    lines = rendered.splitlines()
    assert lines[0] == "| Vùng | Mức lương |"
    assert lines[1] == "| --- | --- |"
    assert lines[2] == "| I | 5.310.000 |"


def test_table_to_markdown_mot_hang_ra_html():
    table = _make_docx_table([["A = B x C"]])
    rendered = tables.table_to_markdown(table)
    assert rendered.startswith("<table>")
    assert "A = B x C" in rendered


def test_is_signature_table_nhan_dien_theo_noi_dung():
    assert tables.is_signature_table(T("| TM. THỦ TƯỚNG | |\n| --- | --- |\n| A | B |"))
    assert tables.is_signature_table(
        T("| Nơi nhận: | |\n| --- | --- |\n| - Như trên | |")
    )


def test_is_signature_table_khong_nhan_dien_bang_thuong():
    assert not tables.is_signature_table(
        T("| Vùng | Mức |\n| --- | --- |\n| I | 5.000 |")
    )


def test_is_signature_table_bo_qua_paragraph():
    # Chỉ block dạng bảng mới được xét, dù nội dung đoạn văn khớp regex.
    assert not tables.is_signature_table(P("TM. THỦ TƯỚNG"))


def test_filter_middle_tables_loai_bang_dinh_kem():
    blocks = [
        P("Điều 1. Phạm vi điều chỉnh"),
        T("FILE ĐƯỢC ĐÍNH KÈM THEO VĂN BẢN"),
        T("| Vùng | Mức |\n| --- | --- |\n| I | 5.000 |"),
    ]
    kept, warnings = tables.filter_middle_tables(blocks)
    assert kept == [blocks[0], blocks[2]]
    assert [w.code for w in warnings] == ["dropped_attachment_table"]


def test_filter_middle_tables_khong_co_bang_dinh_kem_giu_nguyen():
    blocks = [P("Điều 1. A"), T("| Vùng | Mức |\n| --- | --- |\n| I | 5.000 |")]
    kept, warnings = tables.filter_middle_tables(blocks)
    assert kept == blocks
    assert warnings == []


# ==========================================================================
# emitter.py -- không đổi so với bản cũ, chỉ đổi signature (không còn
# refs/footnote_map/preamble_end -- front/back matter đã được cắt từ trước).
# ==========================================================================


def test_emit_anh_xa_heading_co_ban():
    blocks = [
        P("Chương I"),
        P("NHỮNG QUY ĐỊNH CHUNG"),
        P("Điều 1. Phạm vi điều chỉnh"),
        P("1. Nội dung khoản một."),
        P("a) Điểm a của khoản một."),
    ]
    parts = emitter.emit(blocks)
    assert parts[0] == "## Chương I. NHỮNG QUY ĐỊNH CHUNG"
    assert parts[1] == "#### Điều 1. Phạm vi điều chỉnh"
    assert parts[2] == "##### Khoản 1"
    assert parts[3] == "Nội dung khoản một."
    assert parts[4] == "a) Điểm a của khoản một."


def test_emit_phan_ghep_tieu_de_va_khong_o_trong_phu_luc():
    blocks = [P("Phần thứ nhất"), P("QUY ĐỊNH CHUNG"), P("1. Nội dung.")]
    parts = emitter.emit(blocks)
    assert parts[0] == "# Phần thứ nhất. QUY ĐỊNH CHUNG"
    assert parts[1] == "##### Khoản 1"
    assert parts[2] == "Nội dung."


def test_emit_bang_duoc_xuat_nguyen_van():
    table_text = "| a | b |\n| --- | --- |\n| 1 | 2 |"
    blocks = [P("Điều 1. A"), T(table_text)]
    parts = emitter.emit(blocks)
    assert parts[1] == table_text


def test_emit_phu_luc_khoan_ngan_gop_vao_heading():
    blocks = [
        P("PHỤ LỤC"),
        P("DANH MỤC TỈNH THÀNH"),
        P("28. Thành phố Hồ Chí Minh"),
    ]
    parts = emitter.emit(blocks)
    assert parts[0] == "# PHỤ LỤC — DANH MỤC TỈNH THÀNH"
    assert parts[1] == "##### 28. Thành phố Hồ Chí Minh"


def test_emit_phu_luc_khoan_dai_khong_gop_vao_heading():
    long_content = "Nội dung dài " * 20  # > 120 ký tự
    blocks = [P("PHỤ LỤC"), P("BẢNG DANH MỤC"), P(f"1. {long_content}")]
    parts = emitter.emit(blocks)
    assert parts[1] == "##### Khoản 1"
    assert parts[2] == long_content


def test_emit_tach_phan_kem_theo_khoi_tieu_de_phu_luc():
    blocks = [
        P("PHỤ LỤC"),
        P("BẢNG LƯƠNG (Kèm theo Nghị định số 293/2025/NĐ-CP)"),
        P("1. Nội dung."),
    ]
    parts = emitter.emit(blocks)
    assert parts[0] == "# PHỤ LỤC — BẢNG LƯƠNG"
    assert parts[1] == "(Kèm theo Nghị định số 293/2025/NĐ-CP)"


def test_emit_ra_khoi_phu_luc_khi_gap_chuong_khoan_khong_gop_nua():
    blocks = [
        P("PHỤ LỤC"),
        P("DANH MỤC"),
        P("28. Hà Nội"),
        P("Chương I"),
        P("PHẦN MỚI"),
        P("1. Nội dung."),
    ]
    parts = emitter.emit(blocks)
    assert parts[-2] == "##### Khoản 1"
    assert parts[-1] == "Nội dung."


# --- Regression: marker "[n]" dính liền phá RE_KHOAN/RE_DIEU nếu không --
# gỡ trước khi khớp heading (formatting_spec.md mục 1.1, "Lưu ý quan
# trọng"; xem thêm test strip_markers ở patterns.py) --------------------


def test_emit_go_marker_truoc_khi_khop_khoan():
    blocks = [
        P("Điều 1. Phạm vi điều chỉnh"),
        P("1.[2] Bảo hiểm y tế là hình thức bắt buộc."),
    ]
    parts = emitter.emit(blocks)
    assert parts[1] == "##### Khoản 1"
    assert parts[2] == "Bảo hiểm y tế là hình thức bắt buộc."


def test_emit_go_marker_truoc_khi_khop_dieu():
    blocks = [P("Điều 7a.[3] Bảo hiểm xã hội")]
    parts = emitter.emit(blocks)
    assert parts[0] == "#### Điều 7a. Bảo hiểm xã hội"


def test_emit_khong_go_marker_trong_bang():
    # Bảng KHÔNG bị strip -- marker còn sót trong bảng là dấu hiệu bất
    # thường, validator.py cảnh báo riêng (orphan_footnote), emitter không
    # tự ý sửa nội dung bảng.
    table_text = "| a | [2] |\n| --- | --- |\n| 1 | 2 |"
    blocks = [P("Điều 1. A"), T(table_text)]
    parts = emitter.emit(blocks)
    assert parts[1] == table_text


# ==========================================================================
# validator.py -- chỉ chạy trên markdown vùng nội dung ở giữa, không còn
# khối YAML/blockquote chú thích để bóc tách.
# ==========================================================================


def test_validate_khong_co_heading_canh_bao():
    warnings = validator.validate("Một đoạn văn thường không có heading nào.")
    assert any(w.code == "no_heading" for w in warnings)


def test_validate_heading_qua_sau_canh_bao():
    markdown = "#### Điều 1. A\n\n###### quá sâu"
    warnings = validator.validate(markdown)
    assert any(w.code == "heading_too_deep" for w in warnings)


def test_validate_bo_qua_level_skip_hop_le_trong_phu_luc():
    markdown = "# PHỤ LỤC\n\n##### 1. Hà Nội\n\nNội dung."
    warnings = validator.validate(markdown)
    assert not any(w.code == "heading_level_skip" for w in warnings)


def test_validate_bao_level_skip_ngoai_phu_luc():
    markdown = "### Mục 1\n\n##### Khoản 1\n\nNội dung."
    warnings = validator.validate(markdown)
    assert any(w.code == "heading_level_skip" for w in warnings)


def test_validate_dieu_khong_tang_dan():
    markdown = "#### Điều 2. B\n\nNội dung.\n\n#### Điều 1. A\n\nNội dung."
    warnings = validator.validate(markdown)
    assert any(w.code == "dieu_not_monotonic" for w in warnings)


def test_validate_khoan_khong_tang_dan():
    markdown = (
        "#### Điều 1. A\n\n##### Khoản 2\n\nNội dung.\n\n##### Khoản 1\n\nNội dung."
    )
    warnings = validator.validate(markdown)
    assert any(w.code == "khoan_not_monotonic" for w in warnings)


def test_validate_dieu_rong_canh_bao():
    markdown = "#### Điều 1. A\n\n#### Điều 2. B\n\nNội dung điều 2."
    warnings = validator.validate(markdown)
    assert any(w.code == "empty_dieu" and w.detail == "Điều 1. A" for w in warnings)


def test_validate_dieu_co_bang_khong_bi_bao_rong():
    markdown = "#### Điều 1. A\n\n<table>\n<tr><td>x</td></tr>\n</table>"
    warnings = validator.validate(markdown)
    assert not any(w.code == "empty_dieu" for w in warnings)


def test_validate_heading_dai_bat_thuong():
    markdown = "#### Điều 1. " + "A" * 200
    warnings = validator.validate(markdown)
    assert any(w.code == "suspicious_heading_length" for w in warnings)


def test_validate_marker_con_sot_bi_canh_bao_orphan_footnote():
    # emitter._strip_block_markers không strip trong bảng (chỉ paragraph) --
    # nếu marker còn sót ở bất kỳ đâu trong output, đó là bất thường cần rà
    # lại thủ công (formatting_spec.md mục 1.1).
    markdown = "#### Điều 1. A\n\nNội dung còn sót [5] marker."
    warnings = validator.validate(markdown)
    matching = [w for w in warnings if w.code == "orphan_footnote"]
    assert len(matching) == 1
    assert matching[0].detail == "[5]"


def test_validate_nhieu_marker_con_sot_moi_marker_mot_canh_bao():
    markdown = "Nội dung [1] và [2] còn sót."
    warnings = validator.validate(markdown)
    matching = [w for w in warnings if w.code == "orphan_footnote"]
    assert [w.detail for w in matching] == ["[1]", "[2]"]


def test_validate_khong_con_marker_khong_bi_canh_bao_orphan_footnote():
    markdown = "#### Điều 1. A\n\nNội dung sạch không có marker."
    warnings = validator.validate(markdown)
    assert not any(w.code == "orphan_footnote" for w in warnings)


# ==========================================================================
# frontmatter.py
# ==========================================================================


def test_frontmatter_find_boundary_tra_ve_chi_so_dau_tien_la_heading():
    blocks = [P("CHÍNH PHỦ"), P("Điều 1. Phạm vi.")]
    assert frontmatter.find_boundary(blocks) == 1


def test_frontmatter_find_boundary_khong_co_heading_tra_ve_do_dai_list():
    blocks = [P("CHÍNH PHỦ"), P("Số: 1/2025/NĐ-CP")]
    assert frontmatter.find_boundary(blocks) == len(blocks)


def test_frontmatter_find_boundary_bo_qua_dong_danh_so_o_dan_nhap():
    blocks = [
        P("Căn cứ Luật Nhà giáo số 73/2025/QH15..."),
        P("1. Luật Nhà giáo số 73/2025/QH15."),
        P("Chương I"),
        P("NHỮNG QUY ĐỊNH CHUNG"),
    ]
    assert frontmatter.find_boundary(blocks) == 2


def test_frontmatter_find_boundary_chi_xet_paragraph():
    blocks = [T("Điều 1. Trong bảng không phải heading thật.")]
    assert frontmatter.find_boundary(blocks) == len(blocks)


# --------------------------------------------------------------------------
# find_title (mục 1.1 spec, heuristic sửa lại lần 2 -- vị trí tương đối so
# với dòng loại văn bản, không dựa bold/viết hoa của chính dòng tên văn bản).
# --------------------------------------------------------------------------


def test_find_title_tim_thay_ngay_sau_dong_loai_van_ban():
    blocks = [
        P("CHÍNH PHỦ", bold=True),
        P("LUẬT"),
        P("BẢO HIỂM Y TẾ", bold=True),
        P("Căn cứ Hiến pháp..."),
    ]
    assert frontmatter.find_title(blocks) == 2


def test_find_title_dong_ten_khong_bold_van_tim_thay():
    # 2/6 file thật (dạng Nghị định) có dòng tên văn bản hoàn toàn không
    # bold -- heuristic không được loại nó vì thiếu bold (mục 1.1 spec).
    blocks = [
        P("CHÍNH PHỦ", bold=True),
        P("NGHỊ ĐỊNH"),
        P("Quy định mức lương tối thiểu..."),
    ]
    assert frontmatter.find_title(blocks) == 2


def test_find_title_khong_co_dong_loai_van_ban_tra_ve_none():
    blocks = [P("CHÍNH PHỦ", bold=True), P("Số: 1/2025/NĐ-CP")]
    assert frontmatter.find_title(blocks) is None


def test_find_title_dong_loai_van_ban_la_block_cuoi_khong_co_block_sau_tra_ve_none():
    blocks = [P("CHÍNH PHỦ", bold=True), P("LUẬT")]
    assert frontmatter.find_title(blocks) is None


def test_find_title_block_ngay_sau_la_bang_khong_phai_paragraph_tra_ve_none():
    blocks = [P("CHÍNH PHỦ", bold=True), P("LUẬT"), T("Bảng không phải tên văn bản.")]
    assert frontmatter.find_title(blocks) is None


def test_find_title_block_ngay_sau_rong_tra_ve_none():
    blocks = [P("CHÍNH PHỦ", bold=True), P("LUẬT"), P("")]
    assert frontmatter.find_title(blocks) is None


def test_find_title_lay_dong_loai_van_ban_dau_tien_neu_co_nhieu_ung_vien():
    # Chỉ có nghĩa khi văn bản trích dẫn "LUẬT" trong 1 dòng dẫn nhập khác --
    # find_title lấy khớp ĐẦU TIÊN, không phải khớp cuối cùng.
    blocks = [
        P("CHÍNH PHỦ", bold=True),
        P("LUẬT"),
        P("BẢO HIỂM Y TẾ", bold=True),
        P("NGHỊ ĐỊNH"),
        P("Không liên quan."),
    ]
    assert frontmatter.find_title(blocks) == 2


# --------------------------------------------------------------------------
# build_prompts (mục 1.3 spec -- thuần, KHÔNG gọi Groq, chỉ dựng prompt).
# --------------------------------------------------------------------------


def test_frontmatter_build_prompts_rong_tra_ve_jobs_rong():
    jobs = frontmatter.build_prompts([])
    assert jobs.title_text is None
    assert jobs.before_prompts == []
    assert jobs.after_prompts == []


def test_frontmatter_build_prompts_khong_tim_thay_title_toan_bo_vao_before():
    blocks = [P("CHÍNH PHỦ", bold=True), P("Số: 1/2025/NĐ-CP")]
    jobs = frontmatter.build_prompts(blocks)
    assert jobs.title_text is None
    assert len(jobs.before_prompts) == 1
    assert jobs.after_prompts == []
    assert "**CHÍNH PHỦ**" in jobs.before_prompts[0]
    assert "Số: 1/2025/NĐ-CP" in jobs.before_prompts[0]


def test_frontmatter_build_prompts_tim_thay_title_tach_before_va_after():
    blocks = [
        P("CHÍNH PHỦ", bold=True),
        P("LUẬT"),
        P("BẢO HIỂM Y TẾ", bold=True),
        P("Căn cứ Hiến pháp..."),
    ]
    jobs = frontmatter.build_prompts(blocks)
    assert jobs.title_text == "BẢO HIỂM Y TẾ"
    assert len(jobs.before_prompts) == 1
    assert "**CHÍNH PHỦ**" in jobs.before_prompts[0]
    assert "LUẬT" in jobs.before_prompts[0]  # dòng loại văn bản ở lại `before`
    assert "BẢO HIỂM Y TẾ" not in jobs.before_prompts[0]  # tên văn bản không vào prompt
    assert len(jobs.after_prompts) == 1
    assert "Căn cứ Hiến pháp..." in jobs.after_prompts[0]


def test_frontmatter_build_prompts_title_la_block_cuoi_khong_co_after():
    blocks = [P("CHÍNH PHỦ", bold=True), P("LUẬT"), P("BẢO HIỂM Y TẾ", bold=True)]
    jobs = frontmatter.build_prompts(blocks)
    assert jobs.title_text == "BẢO HIỂM Y TẾ"
    assert len(jobs.before_prompts) == 1
    assert jobs.after_prompts == []


def test_frontmatter_build_prompts_chia_nhieu_chunk_khi_vuot_gioi_han(
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.setattr(llm_client, "get_chunk_token_limit", lambda: 1)
    blocks = [P("CHÍNH PHỦ", bold=True), P("Số: 1/2025/NĐ-CP")]
    jobs = frontmatter.build_prompts(blocks)
    assert jobs.title_text is None
    assert len(jobs.before_prompts) == 2


# --------------------------------------------------------------------------
# assemble (mục 1.3 spec -- thuần, KHÔNG gọi Groq, ghép kết quả Groq đã nhận
# từ `pipeline.py`).
# --------------------------------------------------------------------------


def test_frontmatter_assemble_rong_hoan_toan_khong_canh_bao():
    markdown, warnings = frontmatter.assemble(None, [], [])
    assert markdown == ""
    assert warnings == []


def test_frontmatter_assemble_khong_tim_thay_title_van_ghep_phat_canh_bao():
    markdown, warnings = frontmatter.assemble(None, ["**CHÍNH PHỦ**"], [])
    assert markdown == "**CHÍNH PHỦ**"
    assert [w.code for w in warnings] == ["frontmatter_title_not_found"]


def test_frontmatter_assemble_tim_thay_title_chen_heading_giua_before_va_after():
    markdown, warnings = frontmatter.assemble(
        "BẢO HIỂM Y TẾ", ["**CHÍNH PHỦ**"], ["*Căn cứ Hiến pháp...*"]
    )
    assert markdown == "**CHÍNH PHỦ**\n\n# BẢO HIỂM Y TẾ\n\n*Căn cứ Hiến pháp...*"
    assert warnings == []


def test_frontmatter_assemble_before_loi_bo_qua_phan_do_giu_heading():
    markdown, warnings = frontmatter.assemble("BẢO HIỂM Y TẾ", [None], ["Sau"])
    assert markdown == "# BẢO HIỂM Y TẾ\n\nSau"
    assert [w.code for w in warnings] == ["llm_frontmatter_conversion_failed"]


def test_frontmatter_assemble_after_loi_bo_qua_phan_do_giu_heading():
    markdown, warnings = frontmatter.assemble("BẢO HIỂM Y TẾ", ["Trước"], [None])
    assert markdown == "Trước\n\n# BẢO HIỂM Y TẾ"
    assert [w.code for w in warnings] == ["llm_frontmatter_conversion_failed"]


def test_frontmatter_assemble_heading_nguyen_van_khong_qua_groq_du_ca_hai_loi():
    # Groq lỗi cho cả before/after -- dòng heading vẫn được chèn vì
    # `find_title` (chạy trước, không phụ thuộc kết quả Groq) đã xác định
    # `title_text` (mục 1.1 spec: "dòng heading vẫn luôn được chèn nếu tìm
    # thấy, kể cả khi before/after lỗi và bị bỏ qua").
    markdown, warnings = frontmatter.assemble("BẢO HIỂM Y TẾ", [None], [None])
    assert markdown == "# BẢO HIỂM Y TẾ"
    assert [w.code for w in warnings] == [
        "llm_frontmatter_conversion_failed",
        "llm_frontmatter_conversion_failed",
    ]


def test_frontmatter_assemble_khong_co_title_va_loi_ca_hai_canh_bao():
    markdown, warnings = frontmatter.assemble(None, [None], [])
    assert markdown == ""
    assert [w.code for w in warnings] == [
        "llm_frontmatter_conversion_failed",
        "frontmatter_title_not_found",
    ]


def test_frontmatter_assemble_ap_dung_escape_setext_underline():
    # Bug setext heading (mục 1.1 spec) -- markdown ghép xong phải qua
    # `escape_setext_underline` trước khi trả về.
    markdown, _ = frontmatter.assemble(None, ["Tên cơ quan\n-----"], [])
    assert markdown == "Tên cơ quan\n\n-----"


# ==========================================================================
# backmatter.py
# ==========================================================================


def test_backmatter_find_boundary_lay_bang_ky_cuoi_cung():
    blocks = [T("TM. THỦ TƯỚNG"), P("Ở giữa."), T("Nơi nhận:")]
    assert backmatter.find_boundary(blocks) == 2


def test_backmatter_find_boundary_khong_co_bang_ky_tra_ve_none():
    blocks = [P("Nội dung.")]
    assert backmatter.find_boundary(blocks) is None


def test_backmatter_split_backmatter_co_noi_dung_sau_bang_ky():
    blocks = [P("Ở giữa."), T("TM. THỦ TƯỚNG"), P("[1] Ghi chú.")]
    middle, back, warnings = backmatter.split_backmatter(blocks)
    assert middle == [blocks[0]]
    assert back == [blocks[2]]
    assert [w.code for w in warnings] == ["dropped_noi_nhan_table"]


def test_backmatter_split_backmatter_khong_con_block_sau_bang_ky():
    blocks = [P("Ở giữa."), T("TM. THỦ TƯỚNG")]
    middle, back, warnings = backmatter.split_backmatter(blocks)
    assert middle == [blocks[0]]
    assert back == []
    assert [w.code for w in warnings] == ["dropped_noi_nhan_table"]


def test_backmatter_split_backmatter_khong_co_bang_ky():
    blocks = [P("Ở giữa.")]
    middle, back, warnings = backmatter.split_backmatter(blocks)
    assert middle == blocks
    assert back == []
    assert warnings == []


# --------------------------------------------------------------------------
# build_prompts (mục 1.3 spec -- thuần, KHÔNG gọi Groq).
# --------------------------------------------------------------------------


def test_backmatter_build_prompts_rong_tra_ve_danh_sach_rong():
    assert backmatter.build_prompts([]) == []


def test_backmatter_build_prompts_co_noi_dung():
    prompts = backmatter.build_prompts([P("[1] Ghi chú sửa đổi.")])
    assert len(prompts) == 1
    assert "[1] Ghi chú sửa đổi." in prompts[0]


def test_backmatter_build_prompts_giu_nguyen_marker_khong_bi_strip():
    # Khác với vùng nội dung ở giữa (strip_markers), marker "[n]" ở back
    # matter là số thứ tự chú thích thật (vd. "[1] Luật Công nghiệp...") --
    # PHẢI giữ nguyên khi dựng prompt gửi Groq (formatting_spec.md mục 1.1,
    # "Lưu ý quan trọng").
    prompts = backmatter.build_prompts([P("[1] Luật Công nghiệp có hiệu lực từ...")])
    assert "[1] Luật Công nghiệp có hiệu lực từ..." in prompts[0]


def test_backmatter_build_prompts_chia_nhieu_chunk_khi_vuot_gioi_han(
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.setattr(llm_client, "get_chunk_token_limit", lambda: 1)
    blocks = [P("[1] Ghi chú một."), P("[2] Ghi chú hai.")]
    prompts = backmatter.build_prompts(blocks)
    assert len(prompts) == 2


# --------------------------------------------------------------------------
# assemble (mục 1.3 spec -- thuần, KHÔNG gọi Groq).
# --------------------------------------------------------------------------


def test_backmatter_assemble_rong_khong_canh_bao():
    markdown, warnings = backmatter.assemble([])
    assert markdown == ""
    assert warnings == []


def test_backmatter_assemble_thanh_cong_noi_dung_dung_thu_tu():
    markdown, warnings = backmatter.assemble(["[1] Ghi chú một.", "[2] Ghi chú hai."])
    assert markdown == "[1] Ghi chú một.\n\n[2] Ghi chú hai."
    assert warnings == []


def test_backmatter_assemble_mot_chunk_loi_bo_qua_toan_bo_khong_ghep_do_dang():
    markdown, warnings = backmatter.assemble(["[1] Ghi chú một.", None])
    assert markdown == ""
    assert [w.code for w in warnings] == ["llm_backmatter_conversion_failed"]


def test_backmatter_assemble_ap_dung_escape_setext_underline():
    markdown, _ = backmatter.assemble(["Tên cơ quan\n-----"])
    assert markdown == "Tên cơ quan\n\n-----"


# ==========================================================================
# llm_client.py -- Groq (SDK `groq`), mock hoàn toàn (không gọi API thật)
# ==========================================================================


def _fake_groq_response(content: str | None, *, total_tokens: int | None = 100):
    """Dựng response giả có hình dạng giống `groq.types.chat.ChatCompletion`."""
    message = Mock(content=content)
    choice = Mock(message=message)
    usage = Mock(total_tokens=total_tokens) if total_tokens is not None else None
    return Mock(choices=[choice], usage=usage)


def test_client_groq_dung_tham_so_tu_settings(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("GROQ_API_KEY", "fake-key-khong-goi-thuc-te")
    captured: dict[str, object] = {}

    class FakeGroq:
        def __init__(self, *, api_key: str, timeout: float, max_retries: int) -> None:
            captured["api_key"] = api_key
            captured["timeout"] = timeout
            captured["max_retries"] = max_retries

    monkeypatch.setattr(llm_client, "Groq", FakeGroq)
    llm_client._client.cache_clear()
    try:
        llm_client._client()
        assert captured["api_key"] == "fake-key-khong-goi-thuc-te"
        assert captured["timeout"] == 30.0
        assert captured["max_retries"] == 0
    finally:
        llm_client._client.cache_clear()


def test_client_2_groq_dung_key_2_tu_settings(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("GROQ_API_KEY", "fake-key-1")
    monkeypatch.setenv("GROQ_API_KEY_2", "fake-key-2-rieng")
    captured: dict[str, object] = {}

    class FakeGroq:
        def __init__(self, *, api_key: str, timeout: float, max_retries: int) -> None:
            captured["api_key"] = api_key
            captured["timeout"] = timeout
            captured["max_retries"] = max_retries

    monkeypatch.setattr(llm_client, "Groq", FakeGroq)
    llm_client._client_2.cache_clear()
    try:
        llm_client._client_2()
        assert captured["api_key"] == "fake-key-2-rieng"
        assert captured["timeout"] == 30.0
        assert captured["max_retries"] == 0
    finally:
        llm_client._client_2.cache_clear()


# --------------------------------------------------------------------------
# _convert_one -- gọi/retry/rate-limit 1 chunk (mục 5 spec, logic không đổi
# so với bản `convert_to_markdown` cũ -- chỉ đổi tên + nhận `client`/
# `rate_limiter` tường minh qua tham số thay vì tự đọc singleton bên trong,
# để mỗi thread worker (mục 1.3 spec) truyền vào cặp riêng của chính nó).
# --------------------------------------------------------------------------


def test_convert_one_thanh_cong_goi_dung_tham_so(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("GROQ_API_KEY", "fake-key-khong-goi-thuc-te")
    fake_client = Mock()
    fake_client.chat.completions.create.return_value = _fake_groq_response(
        "  Kết quả markdown  "
    )
    limiter = llm_client._SlidingWindowRateLimiter(tpm_limit=8000, rpm_limit=30)

    result = llm_client._convert_one(
        fake_client, "prompt nội dung", limiter, max_retries=2
    )

    assert result == "Kết quả markdown"
    kwargs = fake_client.chat.completions.create.call_args.kwargs
    assert kwargs["model"] == "openai/gpt-oss-120b"
    assert kwargs["messages"] == [{"role": "user", "content": "prompt nội dung"}]


def test_convert_one_loi_roi_thu_lai_thanh_cong(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("GROQ_API_KEY", "fake-key-khong-goi-thuc-te")
    fake_client = Mock()
    fake_client.chat.completions.create.side_effect = [
        RuntimeError("lỗi mạng giả lập"),
        _fake_groq_response("Kết quả sau khi thử lại"),
    ]
    limiter = llm_client._SlidingWindowRateLimiter(tpm_limit=8000, rpm_limit=30)

    result = llm_client._convert_one(fake_client, "prompt", limiter, max_retries=2)

    assert result == "Kết quả sau khi thử lại"
    assert fake_client.chat.completions.create.call_count == 2


def test_convert_one_het_so_lan_thu_tra_ve_none(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("GROQ_API_KEY", "fake-key-khong-goi-thuc-te")
    fake_client = Mock()
    fake_client.chat.completions.create.side_effect = RuntimeError("lỗi giả lập")
    limiter = llm_client._SlidingWindowRateLimiter(tpm_limit=8000, rpm_limit=30)

    result = llm_client._convert_one(fake_client, "prompt", limiter, max_retries=2)

    assert result is None
    assert fake_client.chat.completions.create.call_count == 2


def test_convert_one_ket_qua_rong_bi_coi_la_loi_va_thu_lai(
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.setenv("GROQ_API_KEY", "fake-key-khong-goi-thuc-te")
    fake_client = Mock()
    fake_client.chat.completions.create.side_effect = [
        _fake_groq_response(""),
        _fake_groq_response("Nội dung thật"),
    ]
    limiter = llm_client._SlidingWindowRateLimiter(tpm_limit=8000, rpm_limit=30)

    result = llm_client._convert_one(fake_client, "prompt", limiter, max_retries=2)

    assert result == "Nội dung thật"
    assert fake_client.chat.completions.create.call_count == 2


def test_convert_one_max_retries_duoc_ton_trong(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("GROQ_API_KEY", "fake-key-khong-goi-thuc-te")
    fake_client = Mock()
    fake_client.chat.completions.create.side_effect = RuntimeError("lỗi giả lập")
    limiter = llm_client._SlidingWindowRateLimiter(tpm_limit=8000, rpm_limit=30)

    result = llm_client._convert_one(fake_client, "prompt", limiter, max_retries=1)

    assert result is None
    assert fake_client.chat.completions.create.call_count == 1


def test_convert_one_thieu_groq_api_key_tra_ve_none_khong_raise(
    monkeypatch: pytest.MonkeyPatch,
):
    """Không có `GROQ_API_KEY` (đúng thực trạng CI) -- `LLMSettings()` raise
    `ValidationError` khi đọc `model_name`, `_convert_one` phải bắt và trả
    `None`, không để lộ exception, không gọi `client` (mục 5, 7 spec: lỗi
    Groq không bao giờ chặn pipeline)."""
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.setattr(
        LLMSettings, "model_config", {**LLMSettings.model_config, "env_file": None}
    )
    fake_client = Mock()
    limiter = llm_client._SlidingWindowRateLimiter(tpm_limit=8000, rpm_limit=30)

    result = llm_client._convert_one(fake_client, "prompt", limiter, max_retries=2)

    assert result is None
    fake_client.chat.completions.create.assert_not_called()


def test_convert_one_cap_nhat_rate_limiter_bang_usage_that(
    monkeypatch: pytest.MonkeyPatch,
):
    """Sau khi gọi thành công, rate limiter phải được cập nhật bằng
    `usage.total_tokens` THẬT từ response, không phải số ước lượng heuristic
    (mục 1.2 spec)."""
    monkeypatch.setenv("GROQ_API_KEY", "fake-key-khong-goi-thuc-te")
    fake_client = Mock()
    fake_client.chat.completions.create.return_value = _fake_groq_response(
        "kết quả", total_tokens=4242
    )
    limiter = llm_client._SlidingWindowRateLimiter(tpm_limit=8000, rpm_limit=30)

    llm_client._convert_one(fake_client, "prompt", limiter, max_retries=2)

    assert [tokens for _, tokens in limiter._entries] == [4242]


def test_convert_one_usage_none_dung_uoc_luong_de_cap_nhat_rate_limiter(
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.setenv("GROQ_API_KEY", "fake-key-khong-goi-thuc-te")
    fake_client = Mock()
    fake_client.chat.completions.create.return_value = _fake_groq_response(
        "kết quả", total_tokens=None
    )
    limiter = llm_client._SlidingWindowRateLimiter(tpm_limit=8000, rpm_limit=30)

    llm_client._convert_one(fake_client, "prompt", limiter, max_retries=2)

    assert len(limiter._entries) == 1


# ==========================================================================
# llm_client.py -- get_chunk_token_limit (formatting_spec.md mục 1.2, 4)
# ==========================================================================


def test_get_chunk_token_limit_doc_tu_settings(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("GROQ_API_KEY", "fake-key-khong-goi-thuc-te")
    monkeypatch.setenv("CHUNK_TOKEN_LIMIT", "999")
    assert llm_client.get_chunk_token_limit() == 999


def test_get_chunk_token_limit_fallback_khi_thieu_api_key(
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.setattr(
        LLMSettings, "model_config", {**LLMSettings.model_config, "env_file": None}
    )
    assert llm_client.get_chunk_token_limit() == 1500


# ==========================================================================
# llm_client.py -- _SlidingWindowRateLimiter (formatting_spec.md mục 1.2, 6)
# ==========================================================================


def test_rate_limiter_khong_cho_khi_cua_so_dang_rong():
    limiter = llm_client._SlidingWindowRateLimiter(tpm_limit=8000, rpm_limit=30)
    # Không có gì để chờ hết hạn -- kể cả khi ước lượng vượt ngưỡng an toàn.
    limiter.wait_if_needed(1_000_000)  # không raise, không treo


def test_rate_limiter_cho_toi_khi_entry_cu_nhat_het_han_do_vuot_tpm(
    monkeypatch: pytest.MonkeyPatch,
):
    limiter = llm_client._SlidingWindowRateLimiter(tpm_limit=100, rpm_limit=30)
    now = [0.0]
    monkeypatch.setattr(llm_client.time, "monotonic", lambda: now[0])
    sleep_calls: list[float] = []

    def fake_sleep(seconds: float) -> None:
        sleep_calls.append(seconds)
        now[0] += seconds

    monkeypatch.setattr(llm_client.time, "sleep", fake_sleep)

    limiter.record(now[0], 50)  # 50 token, an toàn TPM = 90
    limiter.wait_if_needed(50)  # 50 + 50 = 100 > 90 -- phải chờ

    assert sleep_calls == [60.0]


def test_rate_limiter_cho_toi_khi_vuot_rpm(monkeypatch: pytest.MonkeyPatch):
    limiter = llm_client._SlidingWindowRateLimiter(tpm_limit=8000, rpm_limit=1)
    now = [0.0]
    monkeypatch.setattr(llm_client.time, "monotonic", lambda: now[0])
    sleep_calls: list[float] = []

    def fake_sleep(seconds: float) -> None:
        sleep_calls.append(seconds)
        now[0] += seconds

    monkeypatch.setattr(llm_client.time, "sleep", fake_sleep)

    limiter.record(now[0], 10)  # 1 request đã ghi nhận, an toàn RPM = 0.9
    limiter.wait_if_needed(10)  # 1 + 1 = 2 > 0.9 -- phải chờ, bất kể token

    assert sleep_calls == [60.0]


def test_rate_limiter_khong_cho_khi_du_du_du_cho_trong_cua_so(
    monkeypatch: pytest.MonkeyPatch,
):
    limiter = llm_client._SlidingWindowRateLimiter(tpm_limit=8000, rpm_limit=30)
    now = [0.0]
    monkeypatch.setattr(llm_client.time, "monotonic", lambda: now[0])
    sleep_calls: list[float] = []
    monkeypatch.setattr(
        llm_client.time, "sleep", lambda seconds: sleep_calls.append(seconds)
    )

    limiter.record(now[0], 100)
    limiter.wait_if_needed(100)  # 100 + 100 << TPM an toàn 7200

    assert sleep_calls == []


def test_rate_limiter_don_entry_het_han_khoi_cua_so(monkeypatch: pytest.MonkeyPatch):
    limiter = llm_client._SlidingWindowRateLimiter(tpm_limit=8000, rpm_limit=30)
    now = [0.0]
    monkeypatch.setattr(llm_client.time, "monotonic", lambda: now[0])

    limiter.record(0.0, 100)
    now[0] = 61.0  # entry cũ hơn 60s -- phải bị dọn khỏi cửa sổ
    limiter._evict_expired(now[0])

    assert list(limiter._entries) == []


# ==========================================================================
# llm_client.py -- convert_chunks_concurrently (formatting_spec.md mục 1.3
# -- dispatch động 2 API key Groq). Gọi qua `_REAL_CONVERT_CHUNKS_CONCURRENTLY`
# (tham chiếu hàm THẬT) để không bị `_default_llm_stub` (autouse) che mất.
# ==========================================================================


def test_convert_chunks_concurrently_rong_tra_ve_rong():
    assert _REAL_CONVERT_CHUNKS_CONCURRENTLY([]) == []


def test_convert_chunks_concurrently_thieu_settings_tra_ve_none_het(
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.setattr(
        LLMSettings, "model_config", {**LLMSettings.model_config, "env_file": None}
    )
    assert _REAL_CONVERT_CHUNKS_CONCURRENTLY(["p1", "p2"]) == [None, None]


def test_convert_chunks_concurrently_khong_co_key2_chay_tuan_tu_khong_spawn_thread(
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.setenv("GROQ_API_KEY", "only-key-1")
    monkeypatch.delenv("GROQ_API_KEY_2", raising=False)
    monkeypatch.setattr(
        LLMSettings, "model_config", {**LLMSettings.model_config, "env_file": None}
    )

    fake_client = Mock()
    fake_client.chat.completions.create.return_value = _fake_groq_response("OK")
    monkeypatch.setattr(llm_client, "_client", lambda: fake_client)

    def _client_2_khong_duoc_goi():
        raise AssertionError("_client_2 không được gọi khi thiếu GROQ_API_KEY_2")

    monkeypatch.setattr(llm_client, "_client_2", _client_2_khong_duoc_goi)

    def _khong_duoc_spawn_thread(*args, **kwargs):
        raise AssertionError(
            "threading.Thread không được spawn khi chỉ có 1 key (mục 1.3 spec)"
        )

    monkeypatch.setattr(llm_client.threading, "Thread", _khong_duoc_spawn_thread)

    results = _REAL_CONVERT_CHUNKS_CONCURRENTLY(["p1", "p2", "p3"])

    assert results == ["OK", "OK", "OK"]
    assert fake_client.chat.completions.create.call_count == 3


def test_convert_chunks_concurrently_co_key2_dung_ca_hai_client(
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.setenv("GROQ_API_KEY", "key-1")
    monkeypatch.setenv("GROQ_API_KEY_2", "key-2")

    fake_client_1 = Mock()
    fake_client_1.chat.completions.create.return_value = _fake_groq_response("KQ-1")
    fake_client_2 = Mock()
    fake_client_2.chat.completions.create.return_value = _fake_groq_response("KQ-2")
    monkeypatch.setattr(llm_client, "_client", lambda: fake_client_1)
    monkeypatch.setattr(llm_client, "_client_2", lambda: fake_client_2)

    results = _REAL_CONVERT_CHUNKS_CONCURRENTLY(["p1", "p2", "p3", "p4"])

    assert len(results) == 4
    assert all(result in ("KQ-1", "KQ-2") for result in results)
    total_calls = (
        fake_client_1.chat.completions.create.call_count
        + fake_client_2.chat.completions.create.call_count
    )
    assert total_calls == 4
    # Cả 2 worker phải thực sự được dùng (mục 7 spec: "xác nhận cả 2 key đều
    # thực sự được dùng") -- không phải chỉ 1 client xử lý hết do race.
    assert fake_client_1.chat.completions.create.call_count >= 1
    assert fake_client_2.chat.completions.create.call_count >= 1


def test_convert_chunks_concurrently_giu_dung_thu_tu_goc_du_hoan_thanh_khac_thu_tu(
    monkeypatch: pytest.MonkeyPatch,
):
    """Job ở chỉ số 0 cố tình xử lý CHẬM hơn (sleep) các job còn lại -- dù nó
    hoàn thành SAU CÙNG, kết quả cuối cùng vẫn phải nằm đúng vị trí gốc (chỉ
    số 0), không bị đẩy ra sau theo thứ tự hoàn thành (mục 1.3 spec: "đúng
    thứ tự gốc -- đánh số theo chỉ số, không theo thứ tự hoàn thành")."""
    monkeypatch.setenv("GROQ_API_KEY", "key-1")
    monkeypatch.setenv("GROQ_API_KEY_2", "key-2")

    prompts = ["prompt-cham-SLOW", "prompt-nhanh-1", "prompt-nhanh-2", "prompt-nhanh-3"]

    def _make_side_effect():
        def _side_effect(*, model: str, messages: list[dict[str, str]]):
            prompt = messages[0]["content"]
            if "SLOW" in prompt:
                time.sleep(0.15)
            return _fake_groq_response(f"KQ::{prompt}")

        return _side_effect

    fake_client_1 = Mock()
    fake_client_1.chat.completions.create.side_effect = _make_side_effect()
    fake_client_2 = Mock()
    fake_client_2.chat.completions.create.side_effect = _make_side_effect()
    monkeypatch.setattr(llm_client, "_client", lambda: fake_client_1)
    monkeypatch.setattr(llm_client, "_client_2", lambda: fake_client_2)

    results = _REAL_CONVERT_CHUNKS_CONCURRENTLY(prompts)

    assert results == [f"KQ::{prompt}" for prompt in prompts]


# ==========================================================================
# pipeline.py -- _compose_markdown (ghép 3 phần)
# ==========================================================================


def test_compose_markdown_day_du_ba_phan():
    markdown = pipeline._compose_markdown("FRONT", "MIDDLE", "BACK")
    assert markdown == "FRONT\n\nMIDDLE\n\n---\n\nBACK\n"


def test_compose_markdown_khong_co_front_matter():
    markdown = pipeline._compose_markdown("", "MIDDLE", "")
    assert markdown == "MIDDLE\n"


def test_compose_markdown_khong_co_back_matter():
    markdown = pipeline._compose_markdown("FRONT", "MIDDLE", "")
    assert markdown == "FRONT\n\nMIDDLE\n"


# ==========================================================================
# pipeline.py -- convert_docx_to_markdown (integration DOCX thật, LLM mock)
# ==========================================================================


def _build_docx(path: Path, items: list[tuple[str, object]]) -> None:
    """DSL nhỏ để dựng DOCX test: ("p", text) | ("p_bold", text) | ("table", rows)."""
    document = Document()
    for kind, payload in items:
        if kind == "p":
            document.add_paragraph(str(payload))
        elif kind == "p_bold":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(str(payload))
            run.bold = True
        elif kind == "table":
            rows = payload
            assert isinstance(rows, list)
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for row_index, row in enumerate(rows):
                for col_index, value in enumerate(row):
                    table.cell(row_index, col_index).text = value
        else:
            raise ValueError(f"kind không hỗ trợ: {kind}")
    document.save(path)


def test_convert_docx_to_markdown_ghep_du_front_middle_back(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    captured = _stub_convert_chunks_concurrently(
        monkeypatch, ["**CHÍNH PHỦ**", "[1] Ghi chú sửa đổi."]
    )
    docx_path = tmp_path / "full.docx"
    _build_docx(
        docx_path,
        [
            ("p_bold", "CHÍNH PHỦ"),
            ("p", "Điều 1. Phạm vi điều chỉnh"),
            ("p", "1. Nội dung khoản một."),
            ("table", [["TM. THỦ TƯỚNG"], ["Nguyễn Văn A"]]),
            ("p", "[1] Ghi chú sửa đổi."),
        ],
    )

    result = convert_docx_to_markdown(docx_path)

    assert result.markdown == (
        "**CHÍNH PHỦ**\n\n"
        "#### Điều 1. Phạm vi điều chỉnh\n\n"
        "##### Khoản 1\n\n"
        "Nội dung khoản một.\n\n"
        "---\n\n"
        "[1] Ghi chú sửa đổi.\n"
    )
    assert "TM. THỦ TƯỚNG" not in result.markdown
    assert len(captured) == 2  # 1 chunk front (before) + 1 chunk back matter
    codes = {w.code for w in result.warnings}
    assert "dropped_noi_nhan_table" in codes
    assert "llm_frontmatter_conversion_failed" not in codes
    assert "llm_backmatter_conversion_failed" not in codes


def test_convert_docx_to_markdown_khong_co_back_matter_khong_goi_llm_hai_lan(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    captured = _stub_convert_chunks_concurrently(monkeypatch, ["FRONT_MD"])
    docx_path = tmp_path / "no_back.docx"
    _build_docx(
        docx_path,
        [
            ("p", "CHÍNH PHỦ"),
            ("p", "Điều 1. Phạm vi điều chỉnh"),
            ("p", "1. Nội dung khoản một."),
        ],
    )

    result = convert_docx_to_markdown(docx_path)

    assert result.markdown == (
        "FRONT_MD\n\n#### Điều 1. Phạm vi điều chỉnh\n\n##### Khoản 1\n\nNội dung khoản một.\n"
    )
    assert len(captured) == 1  # chỉ 1 prompt front matter, không có back matter
    codes = {w.code for w in result.warnings}
    assert "llm_backmatter_conversion_failed" not in codes
    assert "dropped_noi_nhan_table" not in codes


def test_convert_docx_to_markdown_llm_loi_bo_qua_front_matter_khong_fail(
    tmp_path: Path,
):
    # Fixture `_default_llm_stub` đã trả None cho mọi prompt mặc định.
    docx_path = tmp_path / "front_fail.docx"
    _build_docx(
        docx_path,
        [
            ("p", "CHÍNH PHỦ"),
            ("p", "Điều 1. Phạm vi điều chỉnh"),
            ("p", "1. Nội dung."),
        ],
    )

    result = convert_docx_to_markdown(docx_path)

    assert (
        result.markdown
        == "#### Điều 1. Phạm vi điều chỉnh\n\n##### Khoản 1\n\nNội dung.\n"
    )
    assert any(w.code == "llm_frontmatter_conversion_failed" for w in result.warnings)


def test_convert_docx_to_markdown_llm_loi_bo_qua_back_matter_khong_fail(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    _stub_convert_chunks_concurrently(monkeypatch, ["FRONT_MD", None])
    docx_path = tmp_path / "back_fail.docx"
    _build_docx(
        docx_path,
        [
            ("p", "CHÍNH PHỦ"),
            ("p", "Điều 1. Phạm vi điều chỉnh"),
            ("p", "1. Nội dung."),
            ("table", [["TM. THỦ TƯỚNG"], ["Nguyễn Văn A"]]),
            ("p", "[1] Ghi chú."),
        ],
    )

    result = convert_docx_to_markdown(docx_path)

    assert result.markdown == (
        "FRONT_MD\n\n#### Điều 1. Phạm vi điều chỉnh\n\n##### Khoản 1\n\nNội dung.\n"
    )
    assert "---" not in result.markdown
    codes = {w.code for w in result.warnings}
    assert "llm_backmatter_conversion_failed" in codes
    assert "dropped_noi_nhan_table" in codes


def test_convert_docx_to_markdown_khong_co_front_matter_khong_goi_llm(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    _forbid_llm_calls(monkeypatch)
    docx_path = tmp_path / "no_front.docx"
    _build_docx(
        docx_path,
        [
            ("p", "Điều 1. Phạm vi điều chỉnh"),
            ("p", "1. Nội dung."),
        ],
    )

    result = convert_docx_to_markdown(docx_path)

    assert (
        result.markdown
        == "#### Điều 1. Phạm vi điều chỉnh\n\n##### Khoản 1\n\nNội dung.\n"
    )
    assert result.warnings == []


def test_convert_docx_to_markdown_goi_convert_chunks_concurrently_dung_1_lan(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    """Mục 1.3 spec: `pipeline.py` gộp toàn bộ job front/back matter của 1
    file thành 1 lệnh gọi `convert_chunks_concurrently` DUY NHẤT."""
    call_count = 0

    def fake(prompts: list[str]) -> list[str | None]:
        nonlocal call_count
        call_count += 1
        return [None] * len(prompts)

    monkeypatch.setattr(llm_client, "convert_chunks_concurrently", fake)
    docx_path = tmp_path / "single_call.docx"
    _build_docx(
        docx_path,
        [
            ("p_bold", "CHÍNH PHỦ"),
            ("p", "Điều 1. Phạm vi điều chỉnh"),
            ("p", "1. Nội dung khoản một."),
            ("table", [["TM. THỦ TƯỚNG"], ["Nguyễn Văn A"]]),
            ("p", "[1] Ghi chú sửa đổi."),
        ],
    )

    convert_docx_to_markdown(docx_path)

    assert call_count == 1


def test_convert_docx_to_markdown_file_rong_bao_loi_ro_rang(tmp_path: Path):
    docx_path = tmp_path / "rong.docx"
    Document().save(docx_path)

    with pytest.raises(ValueError, match="Không trích xuất được nội dung"):
        convert_docx_to_markdown(docx_path)


def test_convert_docx_to_markdown_deterministic_khi_llm_on_dinh(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    monkeypatch.setattr(
        llm_client,
        "convert_chunks_concurrently",
        lambda prompts: ["FRONT_MD"] * len(prompts),
    )
    docx_path = tmp_path / "det.docx"
    _build_docx(
        docx_path,
        [("p", "CHÍNH PHỦ"), ("p", "Điều 1. A"), ("p", "1. Nội dung.")],
    )

    first = convert_docx_to_markdown(docx_path).markdown
    second = convert_docx_to_markdown(docx_path).markdown
    assert first == second


# ==========================================================================
# pipeline.py -- integration trên corpus thật (data/raw/*.docx), LLM mock
# ==========================================================================


@pytest.mark.skipif(not RAW_FILES, reason="data/raw/*.docx không tồn tại")
def test_data_raw_co_du_6_file():
    assert len(RAW_FILES) == 6


def _heading_lines(markdown: str) -> list[str]:
    return [line for line in markdown.splitlines() if line.startswith("#")]


@pytest.mark.skipif(not RAW_FILES, reason="data/raw/*.docx không tồn tại")
@pytest.mark.parametrize("path", RAW_FILES, ids=lambda p: p.stem)
def test_convert_heading_khop_voi_tham_chieu_data_markdown(path: Path):
    """Mục 7 spec: heading mapping cho phần nội dung ở giữa phải khớp với
    `data/markdown/*.md` (tham chiếu, không phải golden-file chính thức).
    Ở đây LLM bị mock trả `None`, nên phần front/back matter do Groq sinh
    luôn rỗng -- ngoại lệ là dòng heading `# <tên văn bản>` (mục 1.1 spec):
    được `frontmatter.find_title` chèn deterministic, không qua Groq, nên
    vẫn xuất hiện kể cả khi mock trả `None`. Vì vậy `data/markdown/*.md`
    tham chiếu phải được regenerate với thiết kế heading mới (dòng tên văn
    bản không kèm tiền tố loại văn bản, không bold) để khớp."""
    reference_path = REFERENCE_MD_DIR / f"{path.stem}.md"
    if not reference_path.exists():
        pytest.skip(f"Không có tham chiếu cho {path.stem}")

    result = convert_docx_to_markdown(path)
    actual_headings = _heading_lines(result.markdown)
    expected_headings = _heading_lines(reference_path.read_text(encoding="utf-8"))
    assert actual_headings == expected_headings


@pytest.mark.skipif(not RAW_FILES, reason="data/raw/*.docx không tồn tại")
@pytest.mark.parametrize("path", RAW_FILES, ids=lambda p: p.stem)
def test_convert_la_ham_thuan_deterministic(path: Path):
    """NT: cùng input phải ra byte-for-byte cùng output, chạy lại nhiều lần
    (LLM mock trả None ổn định, nên toàn bộ output -- không chỉ vùng giữa --
    deterministic ở đây)."""
    first = convert_docx_to_markdown(path).markdown
    second = convert_docx_to_markdown(path).markdown
    assert first == second


@pytest.mark.skipif(not RAW_FILES, reason="data/raw/*.docx không tồn tại")
def test_convert_khong_co_ma_canh_bao_moi():
    """Mục 7 spec: không phát sinh mã QC warning mới ngoài `QcWarningCode`."""
    unknown: set[str] = set()
    for path in RAW_FILES:
        result = convert_docx_to_markdown(path)
        unknown |= {w.code for w in result.warnings} - KNOWN_WARNING_CODES
    assert not unknown, f"Mã QC warning mới, chưa có trong QcWarningCode: {unknown}"


@pytest.mark.skipif(not RAW_FILES, reason="data/raw/*.docx không tồn tại")
def test_convert_file_khong_ton_tai_bao_loi_ro_rang():
    with pytest.raises(PackageNotFoundError):
        convert_docx_to_markdown(RAW_DIR / "khong_ton_tai.docx")


# ==========================================================================
# pipeline.py -- write_atomic / convert_directory
# ==========================================================================


def test_write_atomic_ghi_dung_noi_dung(tmp_path: Path):
    output_path = tmp_path / "out" / "file.md"
    write_atomic(output_path, "nội dung")
    assert output_path.read_text(encoding="utf-8") == "nội dung"
    # Không để lại file tạm.
    assert list(tmp_path.rglob("*.tmp")) == []


def test_write_atomic_don_dep_file_tam_khi_loi(tmp_path: Path, monkeypatch):
    output_path = tmp_path / "file.md"

    def _boom(_fileno):
        raise OSError("giả lập lỗi ghi đĩa")

    monkeypatch.setattr(pipeline.os, "fsync", _boom)
    with pytest.raises(OSError):
        write_atomic(output_path, "nội dung")
    assert not output_path.exists()
    assert list(tmp_path.glob("*.tmp")) == []


def test_convert_directory_loi_mot_file_khong_chan_ca_batch(tmp_path: Path, capsys):
    raw_dir = tmp_path / "raw"
    out_dir = tmp_path / "out"
    raw_dir.mkdir()

    good_source = RAW_FILES[0] if RAW_FILES else None
    if good_source is None:
        pytest.skip("data/raw/*.docx không tồn tại")
    shutil.copy(good_source, raw_dir / "hop_le.docx")
    (raw_dir / "hong.docx").write_bytes(b"khong phai file docx that")
    (raw_dir / "~$khoa.docx").write_bytes(b"file tam cua Word, phai bi bo qua")

    exit_code = convert_directory(raw_dir, out_dir)

    assert exit_code == 1
    assert (out_dir / "hop_le.md").exists()
    assert not (out_dir / "hong.md").exists()
    assert not (out_dir / "~$khoa.md").exists()

    captured = capsys.readouterr()
    assert "hop_le.docx" in captured.out
    assert "Thành công : 1" in captured.out
    assert "Thất bại   : 1" in captured.out


def test_convert_directory_khong_co_docx_tra_ve_0(tmp_path: Path):
    raw_dir = tmp_path / "raw_rong"
    raw_dir.mkdir()
    assert convert_directory(raw_dir, tmp_path / "out") == 0


def test_convert_directory_khong_tra_ve_file_tam(tmp_path: Path):
    if not RAW_FILES:
        pytest.skip("data/raw/*.docx không tồn tại")
    raw_dir = tmp_path / "raw"
    out_dir = tmp_path / "out"
    raw_dir.mkdir()
    shutil.copy(RAW_FILES[0], raw_dir / RAW_FILES[0].name)

    convert_directory(raw_dir, out_dir)

    assert list(out_dir.rglob("*.tmp")) == []


# ==========================================================================
# tools/format_documents.py -- CLI Typer
# ==========================================================================


def test_cli_chuyen_doi_toan_bo_thu_muc(tmp_path: Path):
    if not RAW_FILES:
        pytest.skip("data/raw/*.docx không tồn tại")
    from tools.format_documents import app

    raw_dir = tmp_path / "raw"
    out_dir = tmp_path / "out"
    raw_dir.mkdir()
    shutil.copy(RAW_FILES[0], raw_dir / RAW_FILES[0].name)

    runner = CliRunner()
    result = runner.invoke(app, ["--raw-dir", str(raw_dir), "--out-dir", str(out_dir)])

    assert result.exit_code == 0
    assert (out_dir / f"{RAW_FILES[0].stem}.md").exists()


def test_cli_thuc_thi_duoc_qua_module_chinh():
    """`python tools/format_documents.py --help` không lỗi (kiểm tra entry point).

    Ép NO_COLOR + COLUMNS rộng: Rich (dùng bởi Typer để render --help) tô màu
    ANSI và tự xuống dòng theo bề rộng terminal, nên trên CI (không phải TTY,
    COLUMNS khác máy dev) "raw-dir" có thể bị mã màu hoặc dấu xuống dòng chen
    vào giữa.
    """
    env = {**os.environ, "NO_COLOR": "1", "COLUMNS": "200"}
    completed = subprocess.run(
        [sys.executable, str(PROJECT_ROOT / "tools" / "format_documents.py"), "--help"],
        cwd=PROJECT_ROOT,
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
        env=env,
    )
    assert completed.returncode == 0
    plain_stdout = _ANSI_ESCAPE_RE.sub("", completed.stdout)
    assert "raw-dir" in plain_stdout
